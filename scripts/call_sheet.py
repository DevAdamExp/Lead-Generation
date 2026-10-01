#!/usr/bin/env python3
"""call_sheet.py — work the phone queue, and record what happened.

Phone-first because that is what the data supports: 82% of leads have a verified
phone, 7% a verified email.

Usage:
    python scripts/call_sheet.py --next 20         # show the next 20 to call
    python scripts/call_sheet.py --session         # interactive: call + log
    python scripts/call_sheet.py --export          # XLSX call sheet
    python scripts/call_sheet.py --stats           # funnel so far

Dispositions (interactive mode):
    1 no answer   2 gatekeeper   3 spoke to owner   4 interested
    5 BOOKED      6 not interested   7 bad number    8 do not contact
    s skip (no record)   q quit
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from backend.config import settings                                  # noqa: E402
from backend.database import SyncSessionLocal, create_tables_sync     # noqa: E402
from backend.models import Disposition                               # noqa: E402
from backend.services.outreach import (                              # noqa: E402
    call_script, next_calls, record, stats,
)

CHOICES = {
    "1": Disposition.NO_ANSWER, "2": Disposition.GATEKEEPER,
    "3": Disposition.SPOKE_OWNER, "4": Disposition.INTERESTED,
    "5": Disposition.BOOKED, "6": Disposition.NOT_INTERESTED,
    "7": Disposition.BAD_NUMBER, "8": Disposition.DO_NOT_CONTACT,
}
BAR = "─" * 74


def show(lead, i: int, total: int) -> dict:
    s = call_script(lead)
    print(f"\n{BAR}")
    print(f"[{i}/{total}]  {s['business']}     score {lead.lead_score}"
          f"   conf {lead.data_confidence}")
    print(f"          {s['phone']}"
          + (f"   ·   {lead.address[:52]}" if lead.address else ""))
    print(BAR)
    print(f"  ASK FOR   {s['gatekeeper']}")
    print(f"  OPEN      {s['opener']}")
    if s["reason"]:
        print(f"  PITCH     {s['reason']}")
    for f in s["facts"][1:]:
        print(f"  ALSO      {f}")
    for p in s["pain_points"]:
        print(f"  PAIN      {p[:88]}")
    print(f"  CLOSE     {s['close']}")
    print(f"  IF STALL  {s['objection']}")
    if s["email_followup"]:
        print(f"  EMAIL     {s['email_followup']}  (follow-up only)")
    if lead.website:
        print(f"  SITE      {lead.website}")
    return s


def session(db, limit: int) -> int:
    leads = next_calls(db, limit=limit)
    if not leads:
        print("Nothing due. Either everything is called, or follow-ups aren't ripe yet.")
        return 0
    print(f"{len(leads)} businesses due. Ctrl-C or 'q' to stop — progress is saved "
          f"after each call.")
    done = 0
    for i, lead in enumerate(leads, 1):
        show(lead, i, len(leads))
        while True:
            try:
                key = input("\n  outcome [1-8, s skip, q quit] > ").strip().lower()
            except (EOFError, KeyboardInterrupt):
                print("\nstopped.")
                return done
            if key == "q":
                print(f"\nstopped after {done} logged.")
                return done
            if key == "s":
                break
            if key in CHOICES:
                note = input("  note (optional) > ").strip()
                record(db, lead, CHOICES[key], notes=note)
                done += 1
                break
            print("  ? use 1-8, s, or q")
    print(f"\nqueue finished — {done} calls logged.")
    return done


def export(db, limit: int) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    leads = next_calls(db, limit=limit)
    wb = Workbook()
    ws = wb.active
    ws.title = "Call Sheet"
    cols = [("#", 5), ("Business", 30), ("Phone", 16), ("Score", 7), ("Ask for", 40),
            ("Open with", 52), ("Pitch", 46), ("Pain points", 46), ("Website", 30),
            ("Outcome", 16), ("Notes", 34)]
    for i, (label, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=label)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="0F172A")
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"

    for r, lead in enumerate(leads, 2):
        s = call_script(lead)
        for i, v in enumerate([
            r - 1, s["business"], s["phone"], lead.lead_score, s["gatekeeper"],
            s["opener"], s["reason"] or "", "; ".join(s["pain_points"]),
            lead.website or "", "", "",
        ], 1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=i in (5, 6, 7, 8))

    out = settings.EXPORTS_DIR / "Call_Sheet.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    g = ap.add_mutually_exclusive_group()
    g.add_argument("--next", type=int, metavar="N", help="preview the next N to call")
    g.add_argument("--session", action="store_true", help="interactive call + log")
    g.add_argument("--export", action="store_true", help="write an XLSX call sheet")
    g.add_argument("--stats", action="store_true", help="funnel so far")
    ap.add_argument("--limit", type=int, default=25, help="queue size (default 25)")
    args = ap.parse_args()

    create_tables_sync()
    db = SyncSessionLocal()
    try:
        if args.stats:
            st = stats(db)
            print(f"attempts logged        {st['attempts']}")
            print(f"businesses contacted   {st['businesses_contacted']}")
            print(f"interested             {st['interested']}")
            print(f"BOOKED                 {st['booked']}   ({st['conversion_pct']}%)")
            if st["by_disposition"]:
                print("\nlatest disposition per business:")
                for k, v in sorted(st["by_disposition"].items(), key=lambda x: -x[1]):
                    print(f"  {k:16} {v}")
            return 0

        if args.export:
            p = export(db, args.limit)
            print(f"call sheet -> {p}")
            return 0

        if args.session:
            return 0 if session(db, args.limit) >= 0 else 1

        n = args.next or args.limit
        leads = next_calls(db, limit=n)
        print(f"{len(leads)} due (phone-verified, not closed out, follow-up ripe)")
        for i, lead in enumerate(leads, 1):
            show(lead, i, len(leads))
        return 0
    finally:
        db.close()


if __name__ == "__main__":
    sys.exit(main())
